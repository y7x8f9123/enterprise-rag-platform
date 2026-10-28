from contextlib import asynccontextmanager
from io import BytesIO
from pathlib import Path
import re
from uuid import uuid4

from docx import Document
from fastapi import FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.concurrency import run_in_threadpool
from pydantic import BaseModel, Field
from pypdf import PdfReader

from chunking import split_text
from vector_store import VectorStore

MAX_FILE_SIZE = 20 * 1024 * 1024
MAX_PREVIEW_CHARACTERS = 100_000
ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}
DOCUMENT_ID_PATTERN = re.compile(
    r"^[0-9a-fA-F]{8}-[0-9a-fA-F]{4}-[1-5][0-9a-fA-F]{3}-"
    r"[89abAB][0-9a-fA-F]{3}-[0-9a-fA-F]{12}\.(pdf|docx|txt)$"
)


class HealthResponse(BaseModel):
    status: str
    service: str


class DocumentResponse(BaseModel):
    filename: str
    file_type: str
    size_bytes: int
    characters: int
    words: int
    text: str
    status: str
    document_id: str
    chunk_count: int
    chunks: list[dict]


class IndexResponse(BaseModel):
    status: str
    document_id: str
    filename: str
    file_type: str
    size_bytes: int
    characters: int
    words: int
    chunk_count: int
    indexed_chunks: int


class PreviewResponse(BaseModel):
    text: str
    truncated: bool


class SearchRequest(BaseModel):
    query: str = Field(min_length=1, max_length=10_000)
    limit: int = Field(default=5, ge=1, le=100)


class SearchResult(BaseModel):
    score: float
    document_id: str
    filename: str
    chunk_index: int
    content: str
    start_char: int
    end_char: int


class SearchResponse(BaseModel):
    query: str
    count: int
    results: list[SearchResult]


class DeleteResponse(BaseModel):
    status: str
    document_id: str


def document_error(code: str, message: str) -> dict[str, str]:
    return {"code": code, "message": message}


@asynccontextmanager
async def lifespan(application: FastAPI):
    # One client owns the Qdrant Local directory for the whole process. Qdrant
    # rejects a second process using the same directory; run uvicorn with one
    # worker, or use Qdrant Server when multi-process access is required.
    store = VectorStore()
    application.state.vector_store = store
    try:
        yield
    finally:
        store.close()


app = FastAPI(
    title="EnterpriseMind AI Service",
    description=(
        "Document processing and local RAG service. Local development only; "
        "do not expose this unauthenticated service to a public network."
    ),
    version="0.4.0",
    lifespan=lifespan,
)


@app.get("/api/health", response_model=HealthResponse)
def health_check():
    return HealthResponse(
        status="UP",
        service="EnterpriseMind AI Service",
    )


def extract_text(filename: str, content: bytes) -> str:
    extension = Path(filename).suffix.lower()

    try:
        if extension == ".txt":
            try:
                return content.decode("utf-8-sig")
            except UnicodeDecodeError:
                raise HTTPException(
                    status_code=422,
                    detail=document_error(
                        "INVALID_TEXT_ENCODING",
                        "TXT file must use UTF-8 encoding",
                    ),
                )

        if extension == ".pdf":
            reader = PdfReader(BytesIO(content))

            if reader.is_encrypted:
                raise HTTPException(
                    status_code=422,
                    detail=document_error(
                        "ENCRYPTED_PDF",
                        "Encrypted PDFs are not supported",
                    ),
                )

            pages = [page.extract_text() or "" for page in reader.pages]
            return "\n\n".join(pages)

        if extension == ".docx":
            document = Document(BytesIO(content))
            paragraphs = [
                paragraph.text
                for paragraph in document.paragraphs
                if paragraph.text.strip()
            ]

            for table in document.tables:
                for row in table.rows:
                    cells = [cell.text for cell in row.cells]
                    paragraphs.append(" | ".join(cells))

            return "\n".join(paragraphs)

    except HTTPException:
        raise
    except Exception:
        raise HTTPException(
            status_code=422,
            detail=document_error(
                "DOCUMENT_PARSE_FAILED",
                "Unable to parse document",
            ),
        )

    raise HTTPException(
        status_code=415,
        detail=document_error(
            "UNSUPPORTED_FILE_FORMAT",
            "Unsupported file format",
        ),
    )


async def read_and_validate_upload(file: UploadFile) -> tuple[str, str, bytes]:
    filename = (file.filename or "").replace("\\", "/").rsplit("/", 1)[-1]
    extension = Path(filename).suffix.lower()

    if extension not in ALLOWED_EXTENSIONS:
        await file.close()
        raise HTTPException(
            status_code=415,
            detail="Only PDF, DOCX and TXT are supported",
        )

    content = await file.read(MAX_FILE_SIZE + 1)
    await file.close()

    if not content:
        raise HTTPException(status_code=400, detail="File is empty")
    if len(content) > MAX_FILE_SIZE:
        raise HTTPException(
            status_code=413,
            detail="File exceeds 20 MB limit",
        )

    return filename, extension, content


def parse_and_chunk(
    filename: str,
    content: bytes,
    document_id: str,
) -> tuple[str, list[dict]]:
    text = extract_text(filename, content)
    if not text.strip():
        extension = Path(filename).suffix.lower()
        if extension == ".pdf":
            message = (
                "No extractable text found. This PDF may be scanned or "
                "image-based and may require OCR."
            )
        else:
            message = "No extractable text found in this document."
        raise HTTPException(
            status_code=422,
            detail=document_error("NO_EXTRACTABLE_TEXT", message),
        )

    chunks = split_text(
        text=text,
        document_id=document_id,
        chunk_size=500,
        chunk_overlap=100,
    )
    return text, chunks


@app.post("/api/documents/parse", response_model=DocumentResponse)
async def parse_document(file: UploadFile = File(...)):
    filename, extension, content = await read_and_validate_upload(file)
    document_id = str(uuid4())
    text, chunks = await run_in_threadpool(
        parse_and_chunk,
        filename,
        content,
        document_id,
    )

    return DocumentResponse(
        filename=filename,
        file_type=extension.lstrip("."),
        size_bytes=len(content),
        characters=len(text),
        words=len(text.split()),
        text=text,
        status="PARSED",
        document_id=document_id,
        chunk_count=len(chunks),
        chunks=chunks,
    )


@app.post("/api/documents/preview", response_model=PreviewResponse)
async def preview_document(file: UploadFile = File(...)):
    filename, _, content = await read_and_validate_upload(file)
    text = await run_in_threadpool(extract_text, filename, content)
    truncated = len(text) > MAX_PREVIEW_CHARACTERS
    return PreviewResponse(
        text=text[:MAX_PREVIEW_CHARACTERS],
        truncated=truncated,
    )


@app.post("/api/documents/index", response_model=IndexResponse)
async def index_document(
    request: Request,
    file: UploadFile = File(...),
    document_id: str | None = Form(default=None),
):
    filename, extension, content = await read_and_validate_upload(file)
    resolved_id = (
        str(uuid4()) if document_id is None else document_id.strip()
    )
    if not resolved_id:
        raise HTTPException(
            status_code=422,
            detail="document_id cannot be blank",
        )
    if len(resolved_id) > 200:
        raise HTTPException(
            status_code=422,
            detail="document_id cannot exceed 200 characters",
        )

    text, chunks = await run_in_threadpool(
        parse_and_chunk,
        filename,
        content,
        resolved_id,
    )

    try:
        indexed = await run_in_threadpool(
            request.app.state.vector_store.replace_document,
            resolved_id,
            filename,
            chunks,
        )
    except ValueError as exception:
        raise HTTPException(status_code=422, detail=str(exception))
    except Exception as exception:
        raise HTTPException(
            status_code=503,
            detail="Document indexing failed; the previous index was preserved when possible",
        ) from exception

    return IndexResponse(
        status="INDEXED",
        document_id=resolved_id,
        filename=filename,
        file_type=extension.lstrip("."),
        size_bytes=len(content),
        characters=len(text),
        words=len(text.split()),
        chunk_count=len(chunks),
        indexed_chunks=indexed,
    )


@app.post("/api/search", response_model=SearchResponse)
async def search_documents(body: SearchRequest, request: Request):
    query = body.query.strip()
    if not query:
        raise HTTPException(status_code=422, detail="query cannot be blank")

    try:
        results = await run_in_threadpool(
            request.app.state.vector_store.search,
            query,
            body.limit,
        )
    except ValueError as exception:
        raise HTTPException(status_code=422, detail=str(exception))
    except Exception as exception:
        raise HTTPException(
            status_code=503,
            detail="Search is temporarily unavailable",
        ) from exception

    return SearchResponse(
        query=query,
        count=len(results),
        results=results,
    )


@app.delete(
    "/api/documents/{document_id}/index",
    response_model=DeleteResponse,
)
async def delete_document_index(document_id: str, request: Request):
    if not DOCUMENT_ID_PATTERN.fullmatch(document_id):
        raise HTTPException(status_code=422, detail="Invalid document_id")

    try:
        await run_in_threadpool(
            request.app.state.vector_store.delete_document,
            document_id,
        )
    except ValueError as exception:
        raise HTTPException(status_code=422, detail=str(exception))
    except Exception as exception:
        raise HTTPException(
            status_code=503,
            detail="Vector deletion was not completed",
        ) from exception

    return DeleteResponse(
        status="DELETED",
        document_id=document_id,
    )
