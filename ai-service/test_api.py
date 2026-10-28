import unittest
from io import BytesIO
from unittest.mock import MagicMock, patch

from docx import Document
from fastapi.testclient import TestClient

import main


class FakeVectorStore:
    instances = []

    def __init__(self):
        self.closed = False
        self.index_calls = []
        self.search_calls = []
        self.delete_calls = []
        self.__class__.instances.append(self)

    def replace_document(self, document_id, filename, chunks):
        self.index_calls.append((document_id, filename, chunks))
        return len(chunks)

    def search(self, query, limit):
        self.search_calls.append((query, limit))
        return [{
            "score": 0.91,
            "document_id": "external-42",
            "filename": "handbook.txt",
            "chunk_index": 0,
            "content": "Annual leave is requested in the HR portal.",
            "start_char": 0,
            "end_char": 47,
        }]

    def delete_document(self, document_id):
        self.delete_calls.append(document_id)

    def close(self):
        self.closed = True


class ApiTests(unittest.TestCase):
    def setUp(self):
        FakeVectorStore.instances.clear()
        self.store_patch = patch.object(main, "VectorStore", FakeVectorStore)
        self.store_patch.start()
        self.client_context = TestClient(main.app)
        self.client = self.client_context.__enter__()

    def tearDown(self):
        self.client_context.__exit__(None, None, None)
        self.store_patch.stop()

    def test_health_and_parse_are_preserved(self):
        health = self.client.get("/api/health")
        self.assertEqual(health.status_code, 200)
        self.assertEqual(health.json()["status"], "UP")

        parsed = self.client.post(
            "/api/documents/parse",
            files={"file": ("note.txt", b"hello world", "text/plain")},
        )
        self.assertEqual(parsed.status_code, 200)
        self.assertEqual(parsed.json()["status"], "PARSED")

    def test_index_accepts_external_document_id(self):
        response = self.client.post(
            "/api/documents/index",
            data={"document_id": "external-42"},
            files={
                "file": (
                    "handbook.txt",
                    b"Annual leave is requested in the HR portal.",
                    "text/plain",
                )
            },
        )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload["document_id"], "external-42")
        self.assertEqual(payload["status"], "INDEXED")
        self.assertEqual(payload["indexed_chunks"], 1)
        self.assertEqual(
            FakeVectorStore.instances[0].index_calls[0][0],
            "external-42",
        )

    def test_search_response_contract(self):
        response = self.client.post(
            "/api/search",
            json={"query": "vacation request", "limit": 3},
        )

        self.assertEqual(response.status_code, 200)
        result = response.json()["results"][0]
        self.assertEqual(
            set(result),
            {
                "score",
                "document_id",
                "filename",
                "chunk_index",
                "content",
                "start_char",
                "end_char",
            },
        )

    def test_blank_external_document_id_is_rejected(self):
        response = self.client.post(
            "/api/documents/index",
            data={"document_id": "   "},
            files={"file": ("note.txt", b"hello", "text/plain")},
        )
        self.assertEqual(response.status_code, 422)

    def test_blank_search_is_rejected(self):
        response = self.client.post(
            "/api/search",
            json={"query": "   "},
        )
        self.assertEqual(response.status_code, 422)

    def test_delete_document_index(self):
        document_id = "550e8400-e29b-41d4-a716-446655440000.txt"
        response = self.client.delete(
            f"/api/documents/{document_id}/index"
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json(), {
            "status": "DELETED",
            "document_id": document_id,
        })
        self.assertEqual(
            FakeVectorStore.instances[0].delete_calls,
            [document_id],
        )

    def test_delete_rejects_path_like_document_id(self):
        response = self.client.delete(
            "/api/documents/not-a-valid-id.txt/index"
        )
        self.assertEqual(response.status_code, 422)

    def test_delete_reports_vector_store_failure(self):
        document_id = "550e8400-e29b-41d4-a716-446655440000.txt"
        store = FakeVectorStore.instances[0]
        with patch.object(
            store,
            "delete_document",
            side_effect=RuntimeError("simulated failure"),
        ):
            response = self.client.delete(
                f"/api/documents/{document_id}/index"
            )

        self.assertEqual(response.status_code, 503)
        self.assertEqual(
            response.json()["detail"],
            "Vector deletion was not completed",
        )

    def test_lifespan_closes_store(self):
        store = FakeVectorStore.instances[0]
        self.assertFalse(store.closed)
        self.client_context.__exit__(None, None, None)
        self.assertTrue(store.closed)
        self.client_context = TestClient(main.app)
        self.client = self.client_context.__enter__()

    def test_txt_preview_and_long_preview_truncation(self):
        response = self.client.post(
            "/api/documents/preview",
            files={"file": ("note.txt", "hello 世界".encode(), "text/plain")},
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json(), {
            "text": "hello 世界",
            "truncated": False,
        })

        long_text = "x" * (main.MAX_PREVIEW_CHARACTERS + 1)
        response = self.client.post(
            "/api/documents/preview",
            files={"file": ("long.txt", long_text.encode(), "text/plain")},
        )
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()["truncated"])
        self.assertEqual(
            len(response.json()["text"]),
            main.MAX_PREVIEW_CHARACTERS,
        )

    def test_pdf_preview_extracts_page_text(self):
        first_page = MagicMock()
        first_page.extract_text.return_value = "First page"
        second_page = MagicMock()
        second_page.extract_text.return_value = "Second page"
        reader = MagicMock()
        reader.is_encrypted = False
        reader.pages = [first_page, second_page]

        with patch.object(main, "PdfReader", return_value=reader):
            response = self.client.post(
                "/api/documents/preview",
                files={"file": ("policy.pdf", b"%PDF-test", "application/pdf")},
            )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["text"], "First page\n\nSecond page")

    def test_pdf_with_text_layer_indexes_even_when_it_contains_images(self):
        page = MagicMock()
        page.extract_text.return_value = "Readable policy text"
        page.images = [MagicMock()]
        reader = MagicMock(is_encrypted=False, pages=[page])

        with patch.object(main, "PdfReader", return_value=reader):
            response = self.client.post(
                "/api/documents/index",
                data={"document_id": "document-with-image.pdf"},
                files={"file": ("policy.pdf", b"%PDF-test", "application/pdf")},
            )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["status"], "INDEXED")
        self.assertEqual(len(FakeVectorStore.instances[0].index_calls), 1)

    def test_image_only_and_whitespace_pdf_return_machine_readable_error(self):
        for extracted_text in (None, "  \n\t  "):
            with self.subTest(extracted_text=extracted_text):
                page = MagicMock()
                page.extract_text.return_value = extracted_text
                reader = MagicMock(is_encrypted=False, pages=[page])

                with patch.object(main, "PdfReader", return_value=reader):
                    response = self.client.post(
                        "/api/documents/index",
                        data={"document_id": "image-only.pdf"},
                        files={"file": ("scan.pdf", b"%PDF-test", "application/pdf")},
                    )

                self.assertEqual(response.status_code, 422)
                self.assertEqual(response.json()["detail"], {
                    "code": "NO_EXTRACTABLE_TEXT",
                    "message": (
                        "No extractable text found. This PDF may be scanned or "
                        "image-based and may require OCR."
                    ),
                })
                self.assertEqual(FakeVectorStore.instances[0].index_calls, [])

    def test_docx_preview_includes_paragraphs_and_tables(self):
        document = Document()
        document.add_paragraph("Policy heading")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Name"
        table.cell(0, 1).text = "Value"
        content = BytesIO()
        document.save(content)

        response = self.client.post(
            "/api/documents/preview",
            files={
                "file": (
                    "policy.docx",
                    content.getvalue(),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertIn("Policy heading", response.json()["text"])
        self.assertIn("Name | Value", response.json()["text"])

    def test_preview_rejects_unsupported_and_unparseable_files(self):
        unsupported = self.client.post(
            "/api/documents/preview",
            files={"file": ("unsafe.exe", b"binary", "application/octet-stream")},
        )
        self.assertEqual(unsupported.status_code, 415)

        broken = self.client.post(
            "/api/documents/preview",
            files={"file": ("broken.pdf", b"not a pdf", "application/pdf")},
        )
        self.assertEqual(broken.status_code, 422)
        self.assertEqual(broken.json()["detail"], {
            "code": "DOCUMENT_PARSE_FAILED",
            "message": "Unable to parse document",
        })


if __name__ == "__main__":
    unittest.main()
